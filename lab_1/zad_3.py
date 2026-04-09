from docx import Document
from docx.shared import Inches
import matplotlib.pyplot as plt
import numpy as np
from io import BytesIO
import soundfile as sf
import scipy.fftpack


data, fs = sf.read('SIN/sin_440Hz.wav', dtype='float32')  

def plotAudio(ax,Signal,Fs,Fsize = 2**8,TimeMargin=[0,0.02]):

    print(Fsize)
    
    ax[0].plot(np.arange(0,Signal.shape[0])/Fs, Signal)

    ax[0].set_title("Sygnał w dziedzinie czasu")
    ax[0].set_xlabel("Czas [s]")
    ax[0].set_ylabel("Amplituda")
    ax[0].set_xlim(TimeMargin)
    ax[0].grid(True)

    yf = scipy.fftpack.fft(Signal, n=Fsize)
    
    yf_db = 20*np.log10( np.abs(yf[:fsize//2]))
    xf = np.arange(0,Fs/2,Fs/fsize)

    plt.plot(xf,yf_db)
    
    ax[1].set_title("Widmo częstotliwościowe")
    ax[1].set_xlabel("Częstotliwość [Hz]")
    ax[1].set_ylabel("Amplituda [dB]")
    ax[1].grid()

    idx_max = np.argmax(yf_db)
    
    peak_amp = yf_db[idx_max] 
    peak_freq = xf[idx_max]

    return peak_amp, peak_freq
    

document = Document()
document.add_heading('Zmień ten tytuł',0) 
files=['sin_60Hz.wav','sin_440Hz.wav','sin_8000Hz.wav']
Margins=[[0,0.02],[0.133,0.155]]
fsizes=[2**8,2**12,2**16]

for file in files:
    document.add_heading('Plik - {}'.format(file),2)
    for i,fsize in enumerate(fsizes):
        document.add_heading('Time margin {}'.format([0,0.02]),3) # nagłówek sekcji, mozę być poziom wyżej
        fig ,axs = plt.subplots(2,1,figsize=(10,7)) # tworzenie plota
    
        ############################################################
        # Tu wykonujesz jakieś funkcje i rysujesz wykresy
        ############################################################
        
        data, fs = sf.read(file, dtype=np.int32)


        peak_amp,peak_freq = plotAudio(axs ,data ,fs ,fsize ,[0,0.02])
        

        fig.suptitle('Time margin {}'.format([0,0.02])) # Tytuł wykresu
        fig.tight_layout(pad=1.5) # poprawa czytelności 
        memfile = BytesIO() # tworzenie bufora
        fig.savefig(memfile) # z zapis do bufora 
        
    
        document.add_picture(memfile, width=Inches(6)) # dodanie obrazu z bufora do pliku
        
        memfile.close()
        ############################################################
        # Tu dodajesz dane tekstowe - wartosci, wyjscie funkcji ect.
        document.add_paragraph('wartość losowa = {}'.format(np.random.rand(1))) 
        document.add_paragraph(
            f'Najwyższa wartość w widmie znajduje się dla częstotliwości: {peak_freq:.2f} Hz '
            f'(Amplituda: {peak_amp:.2f} dB).'
        )
        ############################################################

document.save('report.docx') # zapis do pliku



