import numpy as np
import matplotlib.pyplot as plt
import pandas as pd
import cv2
import os
from docx import Document
from docx.shared import Inches
from io import BytesIO


##########################################
### Settings #############################
##########################################

Test = False
Scaling_test = False # run only artificial test for scaling methods

ScalesUp = [5] # list of parameters values
ScalesDown =[0.5] # list of parameters values

OutputRaportFile = ".docx" 

##########################################
### Data Set #############################
##########################################

ImgDir = r'.' # Address of folder with files (do nor delete `r``)


SmallImages=['IMG_SMALL/SMALL_0001.tif','IMG_SMALL/SMALL_0002.png','IMG_SMALL/SMALL_0003.png'
             ,'IMG_SMALL/SMALL_0004.jpg','IMG_SMALL/SMALL_0005.jpg','IMG_SMALL/SMALL_0006.jpg'
             ,'IMG_SMALL/SMALL_0007.jpg','IMG_SMALL/SMALL_0008.jpg','IMG_SMALL/SMALL_0009.jpg',
             'IMG_SMALL/SMALL_0010.jpg'] # list of file names

BigImages=[ #list of dictionaries
    {
        "Filename":"IMG_BIG/BIG_0001.jpg", # File name
        "ROIs":[[10,10,250,250]] # list of Region of interests for this image more then 1 per file
    },
    {
        "Filename":"IMG_BIG/BIG_0002.jpg", # File name
        "ROIs":[[0,0,250,250]] # list of Region of interests for this image more then 1 per file
    },
    {
        "Filename":"IMG_BIG/BIG_0003.jpg", # File name
        "ROIs":[[0,0,250,250]] # list of Region of interests for this image more then 1 per file
    },
    {
        "Filename":"IMG_BIG/BIG_0004.png", # File name
        "ROIs":[[0,0,250,250]] # list of Region of interests for this image more then 1 per file
    }
]


##########################################
### Functions to  ########################
##########################################

# Scaling methods

def NearestNeigbourScaling(In_img, scale):
    height = In_img.shape[0]
    width = In_img.shape[1]

    new_height = int(np.ceil(height * scale))
    new_width = int(np.ceil(width * scale))

    if len(In_img.shape) < 3:
        Out_img = np.zeros((new_height, new_width), dtype=In_img.dtype)
    else:
        Out_img = np.zeros((new_height, new_width, In_img.shape[2]), dtype=In_img.dtype)

    Y_points = np.linspace(0, height - 1, new_height)
    X_points = np.linspace(0, width - 1, new_width)

    for iX, x_val in enumerate(X_points):
        for iY, y_val in enumerate(Y_points):
            
            src_x = int(np.round(x_val))
            src_y = int(np.round(y_val))
            
            Out_img[iY, iX] = In_img[src_y, src_x]

    return Out_img

def BilinearScaling(In_img,scale):
    Out_img=In_img
    ####

    height = In_img.shape[0]
    width = In_img.shape[1]

    new_width = int(np.round(width * scale))
    new_height = int(np.round(height * scale))

    if len(In_img.shape) == 3:
        Out_img = np.zeros((new_height, new_width, In_img.shape[2]), dtype=np.float32)
    else:
        Out_img = np.zeros((new_height, new_width), dtype=np.float32)

    Y_points = np.linspace(0, height - 1, new_height)
    X_points = np.linspace(0, width - 1, new_width)
    
    for iX, x_val in enumerate(X_points):
        for iY, y_val in enumerate(Y_points):
            y_weight = y_val % 1
            x_weight = x_val % 1

            y1 = int(np.floor(y_val))
            x1 = int(np.floor(x_val))
            y2 = min(y1 + 1, height - 1)
            x2 = min(x1 + 1, width - 1)

            f1 = In_img[y1,x1]
            f2 = In_img[y1,x2]
            f3 = In_img[y2,x1]
            f4 = In_img[y2,x2]

            interpolation = (
                f1 * (1 - x_weight) * (1 - y_weight) +
                f2 * x_weight * (1 - y_weight) +
                f3 * (1 - x_weight) * y_weight +
                f4 * x_weight * y_weight
            )

            Out_img[iY, iX] = interpolation


    return Out_img

# Shrinking methods
 
def MeanResizing(In_img,scale):
    Out_img=In_img
    
    height = In_img.shape[0]
    width = In_img.shape[1]

    new_width = int(np.round(width * scale))
    new_height = int(np.round(height * scale))

    if len(In_img.shape) == 3:
        Out_img = np.zeros((new_height, new_width, In_img.shape[2]), dtype=np.float32)
    else:
        Out_img = np.zeros((new_height, new_width), dtype=np.float32)

    Y_points = np.linspace(0, height - 1, new_height)
    X_points = np.linspace(0, width - 1, new_width)
            
    for iX, x_val in enumerate(X_points):
        for iY, y_val in enumerate(Y_points):
            
            ix = np.round(x_val) + np.arange(-3, 4)
            iy = np.round(y_val) + np.arange(-3, 4)

            ix = np.clip(ix, 0, width - 1).astype(int)
            iy = np.clip(iy,0,height - 1).astype(int)
            
            fragment = In_img[iy[0]:iy[-1], ix[0]:ix[-1]]

            fragment_mean = np.mean(fragment)

            Out_img[iY, iX] = fragment_mean

      
    ####
    Out_img = np.clip(Out_img, 0, 255).astype(np.uint8)
    return Out_img



def WeightedMeanResizing(In_img, scale):
    height = In_img.shape[0]
    width = In_img.shape[1]

    new_width = int(np.round(width * scale))
    new_height = int(np.round(height * scale))

    if len(In_img.shape) == 3:
        Out_img = np.zeros((new_height, new_width, In_img.shape[2]), dtype=np.float32)
    else:
        Out_img = np.zeros((new_height, new_width), dtype=np.float32)

    Y_points = np.linspace(0, height - 1, new_height)
    X_points = np.linspace(0, width - 1, new_width)

    # Tworzymy macierz wag 7x7 (ponieważ zakres to od -3 do +3 pikseli)
    # Największe wagi są w centrum, maleją ku brzegom (na wzór rozkładu Gaussa)
    base_weights = np.array([
        [1, 2, 3, 4, 3, 2, 1],
        [2, 4, 6, 8, 6, 4, 2],
        [3, 6, 9, 12, 9, 6, 3],
        [4, 8, 12, 16, 12, 8, 4],
        [3, 6, 9, 12, 9, 6, 3],
        [2, 4, 6, 8, 6, 4, 2],
        [1, 2, 3, 4, 3, 2, 1]
    ], dtype=np.float32)
            
    for iX, x_val in enumerate(X_points):
        for iY, y_val in enumerate(Y_points):
            
            # Wyliczamy "surowe" współrzędne otoczenia
            raw_ix = np.round(x_val) + np.arange(-3, 4)
            raw_iy = np.round(y_val) + np.arange(-3, 4)

            # Sprawdzamy, które indeksy faktycznie mieszczą się w obrazie (zabezpieczenie krawędzi)
            valid_x = (raw_ix >= 0) & (raw_ix < width)
            valid_y = (raw_iy >= 0) & (raw_iy < height)

            # Zostawiamy tylko poprawne współrzędne
            ix = raw_ix[valid_x].astype(int)
            iy = raw_iy[valid_y].astype(int)
            
            # Pobieramy fragment obrazu (dodajemy +1 na końcu, bo w Pythonie wycinanie nie łapie ostatniego elementu)
            fragment = In_img[iy[0]:iy[-1]+1, ix[0]:ix[-1]+1]

            # Wycinamy tylko te wagi, które odpowiadają pikselom mieszczącym się w obrazie
            w = base_weights[valid_y, :][:, valid_x]

            if len(In_img.shape) == 3:
                # Dla RGB musimy rozszerzyć wagi, by przez pomnożenie nałożyć je na 3 warstwy (R, G, B)
                w_3d = w[:, :, np.newaxis]
                # Mnożymy piksele przez wagi, sumujemy po X i Y (axis 0 i 1), a następnie dzielimy przez sumę UŻYTYCH wag
                fragment_mean = np.sum(fragment * w_3d, axis=(0, 1)) / np.sum(w)
            else:
                # Dla obrazu w skali szarości mnożymy normalnie
                fragment_mean = np.sum(fragment * w) / np.sum(w)

            Out_img[iY, iX] = fragment_mean
            
    # Na koniec zawsze obcinamy zera i rzutujemy na standardowy obrazek 8-bit
    Out_img = np.clip(Out_img, 0, 255).astype(np.uint8)
    return Out_img

def MedianResizing(In_img,scale):

        Out_img=In_img
    
        height = In_img.shape[0]
        width = In_img.shape[1]

        new_width = int(np.round(width * scale))
        new_height = int(np.round(height * scale))

        if len(In_img.shape) == 3:
            Out_img = np.zeros((new_height, new_width, In_img.shape[2]), dtype=np.float32)
        else:
            Out_img = np.zeros((new_height, new_width), dtype=np.float32)

        Y_points = np.linspace(0, height - 1, new_height)
        X_points = np.linspace(0, width - 1, new_width)
                
        for iX, x_val in enumerate(X_points):
            for iY, y_val in enumerate(Y_points):
                
                ix = np.round(x_val) + np.arange(-3, 4)
                iy = np.round(y_val) + np.arange(-3, 4)

                ix = np.clip(ix, 0, width - 1).astype(int)
                iy = np.clip(iy,0,height - 1).astype(int)
                
                fragment = In_img[iy[0]:iy[-1], ix[0]:ix[-1]]

                fragment_mean = np.median(fragment)

                Out_img[iY, iX] = fragment_mean
            
        Out_img = np.clip(Out_img, 0, 255).astype(np.uint8)
        return Out_img


def EdgeDetection(img):
    # Create a copy so we don't modify the original image array
    processed_img = img.copy()
    
    # Check if the image is not uint8
    if processed_img.dtype != np.uint8:
        # If the image is normalized between 0.0 and 1.0, scale it to 0-255
        if processed_img.max() <= 1.0:
            processed_img = processed_img * 255
            
        # Clip values to ensure they stay within 0-255 bounds, then convert
        processed_img = np.clip(processed_img, 0, 255).astype(np.uint8)

    # Now Canny will accept it
    edges = cv2.Canny(processed_img, 100, 200)
 
    plt.subplot(121),plt.imshow(img, cmap='gray')
    plt.title('Original Image'), plt.xticks([]), plt.yticks([])
    plt.subplot(122),plt.imshow(edges, cmap='gray')
    plt.title('Edge Image'), plt.xticks([]), plt.yticks([])

    return edges

##########################################
### Main Program  ########################
##########################################

def plot_resize(img, scale, nnscale, bscale, ed_img, ed_nnscale, ed_bscale, mr_img, wmr_img, mdr_img, ed_mr_img, ed_wmr_img,ed_mdr_img, oROI, filename,counter,figsize=(5,5)):
    ROI=(np.array(oROI)*scale).astype(int)
    f,axs=plt.subplots(4,3,num=counter,figsize=figsize) 
    f.suptitle(f"{filename} ROI: {ROI}")
    axs[0,0].imshow(img[oROI[1]:oROI[1]+oROI[3],oROI[0]:oROI[0]+oROI[2],:])
    axs[0,0].set_title("Original")
    axs[0,0].set_axis_off()

    axs[0,1].imshow(nnscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
    axs[0,1].set_title(f"NN scale {scale}")
    axs[0,1].set_axis_off()


    axs[0,2].imshow(bscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
    axs[0,2].set_title(f"Blinear scale {scale}")
    axs[0,2].set_axis_off()

    if len(ed_img.shape)==3:
        axs[1,0].imshow(ed_img[oROI[1]:oROI[1]+oROI[3],oROI[0]:oROI[0]+oROI[2],:])
        axs[1,0].set_title("Edges Original")
        axs[1,0].set_axis_off()
    else:
        axs[1,0].imshow(ed_img[oROI[1]:oROI[1]+oROI[3],oROI[0]:oROI[0]+oROI[2]])
        axs[1,0].set_title("Edges Original")
        axs[1,0].set_axis_off()

    if len(ed_nnscale.shape)==3:
        axs[1,1].imshow(ed_nnscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
        axs[1,1].set_title("Edges NN")
        axs[1,1].set_axis_off()
    else:
        axs[1,1].imshow(ed_nnscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2]])
        axs[1,1].set_title("Edges NN")
        axs[1,1].set_axis_off()
        
    if len(ed_bscale.shape)==3:
        axs[1,2].imshow(ed_bscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
        axs[1,2].set_title("Edges Bilinear")
        axs[1,2].set_axis_off()
    else:
        axs[1,2].imshow(ed_bscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2]])
        axs[1,2].set_title("Edges Bilinear")
        axs[1,2].set_axis_off()

    axs[2,0].imshow(mr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
    axs[2,0].set_title(f"Mean Resizing scale {scale}")
    axs[2,0].set_axis_off()

    axs[2,1].imshow(wmr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
    axs[2,1].set_title(f"Weighted Mean Resizing scale {scale}")
    axs[2,1].set_axis_off()

    axs[2,2].imshow(mdr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
    axs[2,2].set_title(f"Median Resizing scale {scale}")
    axs[2,2].set_axis_off()
        
    if len(ed_mr_img.shape)==3:
        axs[3,0].imshow(ed_mr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
        axs[3,0].set_title("Edges Mean Resizing")
        axs[3,0].set_axis_off()
    else:
        axs[3,0].imshow(ed_mr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2]])
        axs[3,0].set_title("Edges Mean Resizing")
        axs[3,0].set_axis_off()
        
    if len(ed_wmr_img.shape)==3:
        axs[3,1].imshow(ed_wmr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
        axs[3,1].set_title("Edges Weighted Mean Resizing")
        axs[3,1].set_axis_off()
    else:
        axs[3,1].imshow(ed_nnscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2]])
        axs[3,1].set_title("Edges Weighted Mean Resizing")
        axs[3,1].set_axis_off()
        
    if len(ed_mdr_img.shape)==3:
        axs[3,2].imshow(ed_mdr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
        axs[3,2].set_title("Edges Median Resizing")
        axs[3,2].set_axis_off()
    else:
        axs[3,2].imshow(ed_mdr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2]])
        axs[3,2].set_title("Edges Median Resizing")
        axs[3,2].set_axis_off()
        
    return f

def plot_scaling(img, scale, nnscale, bscale, counter, ed_img, ed_nnscale, ed_bscale, file,figsize=(5,5)):
    f,axs=plt.subplots(2,3,num=counter,figsize=figsize) 
    f.suptitle(f"{file}")

    axs[0,0].imshow(img)
    axs[0,0].set_title("Original")
    axs[0,0].set_axis_off()

    axs[0,1].imshow(nnscale)
    axs[0,1].set_title(f"NN scale {scale}")
    axs[0,1].set_axis_off()

    axs[0,2].imshow(bscale)
    axs[0,2].set_title(f"Blinear scale {scale}")
    axs[0,2].set_axis_off()

    axs[1,0].imshow(ed_img)
    axs[1,0].set_title("Edges Original")
    axs[1,0].set_axis_off()

    axs[1,1].imshow(ed_nnscale)
    axs[1,1].set_title("Edges NN")
    axs[1,1].set_axis_off()

    axs[1,2].imshow(ed_bscale)
    axs[1,2].set_title("Edges Bilinear")
    axs[1,2].set_axis_off()
    return f

if Test:
    # test case
    if Scaling_test:
        img= np.zeros((3,3,3),dtype=np.float32)
        img[1,1,:]=1.0
        for scale in ScalesUp:
            f,axs=plt.subplots(1,2)
            nnscale=NearestNeigbourScaling(img,scale)
            axs[0].imshow(nnscale)
            bscale=BilinearScaling(img,scale)
            axs[1].imshow(bscale)

    else:
        counter=1
        for scale in ScalesUp:
            img=plt.imread(os.path.join(ImgDir,SmallImages[0]))
            nnscale=NearestNeigbourScaling(img,scale)
            bscale=BilinearScaling(img,scale)
            ed_img=EdgeDetection(img)
            ed_nnscale=EdgeDetection(nnscale)
            ed_bscale=EdgeDetection(bscale)

            f = plot_scaling(img, scale, nnscale, bscale, counter, ed_img, ed_nnscale, ed_bscale, SmallImages[0])

            counter+=1

        for scale in ScalesDown:    
            img=plt.imread(os.path.join(ImgDir,BigImages[0]["Filename"]))

            nnscale=NearestNeigbourScaling(img,scale)
            bscale=BilinearScaling(img,scale)

            ed_img=EdgeDetection(img)
            ed_nnscale=EdgeDetection(nnscale)
            ed_bscale=EdgeDetection(bscale)

            mr_img=MeanResizing(img,scale)
            wmr_img=WeightedMeanResizing(img,scale)
            mdr_img=MedianResizing(img,scale)

            ed_mr_img=EdgeDetection(mr_img)
            ed_wmr_img=EdgeDetection(wmr_img)
            ed_mdr_img=EdgeDetection(mdr_img)
  
            f= plot_resize(img, scale, nnscale, bscale, ed_img, ed_nnscale, ed_bscale, mr_img, wmr_img, mdr_img, ed_mr_img, ed_wmr_img,ed_mdr_img, BigImages[0]["ROIs"][0], BigImages[0]["Filename"],counter=counter)
            counter+=1
        
    plt.savefig('test.png')

else: 
    # generate raport
    document = Document()
    document.add_heading('Report',0) # tworzenie nagłówków druga wartość to poziom nagłówka 
    document.add_paragraph("Autor: ")
    document.add_paragraph("Proszę wstawić mi 2 jeżeli tego nie wyedytuję")
    document.add_section()
    document.add_heading("Test algorytmów powiększania",1)
    counter = 1 
    for file in SmallImages:
        img=plt.imread(os.path.join(ImgDir,file))
        for scale in ScalesUp:
            nnscale=NearestNeigbourScaling(img,scale)
            bscale=BilinearScaling(img,scale)
            ed_img=EdgeDetection(img)
            ed_nnscale=EdgeDetection(nnscale)
            ed_bscale=EdgeDetection(bscale)
            
            f = plot_scaling(img, scale, nnscale, bscale, counter, ed_img, ed_nnscale, ed_bscale, file) # set figszie

            memfile = BytesIO() 
            f.savefig(memfile)
            document.add_picture(memfile, width=Inches(6)) # set document size
            memfile.close()
            plt.close(f) # <-- Change this from f.clf()
            counter += 1
    document.add_section()
    document.add_heading("Test algorytmów pomniejszania",1)
    for file_dict in BigImages:
        filename=file_dict['Filename']
        img=plt.imread(os.path.join(ImgDir,filename))
        for scale in ScalesDown:
            nnscale=NearestNeigbourScaling(img,scale)
            bscale=BilinearScaling(img,scale)

            ed_img=EdgeDetection(img)
            ed_nnscale=EdgeDetection(nnscale)
            ed_bscale=EdgeDetection(bscale)

            mr_img=MeanResizing(img,scale)
            wmr_img=WeightedMeanResizing(img,scale)
            mdr_img=MedianResizing(img,scale)

            ed_mr_img=EdgeDetection(mr_img)
            ed_wmr_img=EdgeDetection(wmr_img)
            ed_mdr_img=EdgeDetection(mdr_img)

            for ROI in file_dict['ROIs']:

                f = plot_resize(img, scale, nnscale, bscale, ed_img, ed_nnscale, ed_bscale, mr_img, wmr_img, mdr_img, ed_mr_img, ed_wmr_img,ed_mdr_img, ROI, filename,counter=counter) # set figszie

                memfile = BytesIO() 
                f.savefig(memfile)
                document.add_picture(memfile, width=Inches(6)) 
                memfile.close()
                plt.close(f)
                counter += 1
    document.add_section()
    document.add_heading("Podsumowanie i wnioski",1)
    document.add_paragraph("Tu proszę zebrać wszystkie obserwacje na podstawie powyższych wykresów i napisać wnioski.")
    document.save(OutputRaportFile) 
