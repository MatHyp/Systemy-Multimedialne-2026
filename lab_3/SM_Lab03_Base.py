import numpy as np
import matplotlib.pyplot as plt
import pandas as pd
import cv2
import os
from docx import Document
from docx.shared import Inches
from io import BytesIO


##########################################
### Settings #############################
##########################################

Test = False
Scaling_test = False # run only artificial test for scaling methods

ScalesUp = [2, 3, 5] # list of parameters values
ScalesDown = [0.5, 0.1, 0.05] # list of parameters values

OutputRaportFile = "rapo.docx" 

##########################################
### Data Set #############################
##########################################

ImgDir = r'.' # Address of folder with files (do nor delete `r``)


SmallImages=['IMG_SMALL/SMALL_0001.tif','IMG_SMALL/SMALL_0002.png','IMG_SMALL/SMALL_0003.png'
             ,'IMG_SMALL/SMALL_0004.jpg','IMG_SMALL/SMALL_0005.jpg','IMG_SMALL/SMALL_0006.jpg'
             ,'IMG_SMALL/SMALL_0007.jpg','IMG_SMALL/SMALL_0008.jpg','IMG_SMALL/SMALL_0009.jpg',
             'IMG_SMALL/SMALL_0010.jpg'] # list of file names

BigImages=[ #list of dictionaries
    {
        "Filename":"IMG_BIG/BIG_0001.jpg", # File name
        "ROIs": [
            [0,   0,   250, 250],
            [400, 300, 250, 250],
            [700, 100, 250, 250]
        ]
    },
    {
        "Filename":"IMG_BIG/BIG_0002.jpg", # File name
        "ROIs":[[10,  10,  300, 300],   
            [500, 200, 300, 300],   
            [800, 400, 300, 300]]  
    },
    {
        "Filename":"IMG_BIG/BIG_0003.jpg",
        "ROIs":[
            [1750, 2000, 300, 300],  # ROI 1: Środek - Przejście między ciemną górą wieży a jasną cegłą. Mnóstwo detali i wyraźnych krawędzi.
            [2600, 2800, 300, 300],  # ROI 2: Prawa strona - Fasada budynku po prawej. Wyłapie równe, geometryczne rzędy okien.
            [1500, 4800, 300, 300]   # ROI 3: Dół - Kute, czarne ogrodzenie na tle trawy. Super ostre, kontrastowe krawędzie do detekcji (Canny będzie tu świetnie działał).
        ]
    },
    {
        "Filename":"IMG_BIG/BIG_0004.png", # File name
        "ROIs":[[10,  10,  300, 300],   # lewy górny róg
            [500, 200, 300, 300],   # środek
            [800, 400, 300, 300]]  # list of Region of interests for this image more then 1 per file
    },
]


##########################################
### Functions to  ########################
##########################################

# Scaling methods

def NearestNeigbourScaling(In_img, scale):
    height = In_img.shape[0]
    width = In_img.shape[1]

    new_height = int(np.ceil(height * scale))
    new_width = int(np.ceil(width * scale))

    if len(In_img.shape) < 3:
        Out_img = np.zeros((new_height, new_width), dtype=In_img.dtype)
    else:
        Out_img = np.zeros((new_height, new_width, In_img.shape[2]), dtype=In_img.dtype)

    Y_points = np.linspace(0, height - 1, new_height)
    X_points = np.linspace(0, width - 1, new_width)

    for iX, x_val in enumerate(X_points):
        for iY, y_val in enumerate(Y_points):
            
            src_x = int(np.round(x_val))
            src_y = int(np.round(y_val))
            
            Out_img[iY, iX] = In_img[src_y, src_x]

    return Out_img

def BilinearScaling(In_img,scale):
    Out_img=In_img

    if In_img.max() <= 1.0:
        In_img = In_img * 255.0


    height = In_img.shape[0]
    width = In_img.shape[1]

    new_width = int(np.round(width * scale))
    new_height = int(np.round(height * scale))

    if len(In_img.shape) == 3:
        Out_img = np.zeros((new_height, new_width, In_img.shape[2]), dtype=np.float32)
    else:
        Out_img = np.zeros((new_height, new_width), dtype=np.float32)

    Y_points = np.linspace(0, height - 1, new_height)
    X_points = np.linspace(0, width - 1, new_width)
    
    for iX, x_val in enumerate(X_points):
        for iY, y_val in enumerate(Y_points):
            y_weight = y_val % 1
            x_weight = x_val % 1

            y1 = int(np.floor(y_val))
            x1 = int(np.floor(x_val))
            y2 = min(y1 + 1, height - 1)
            x2 = min(x1 + 1, width - 1)

            f1 = In_img[y1,x1]
            f2 = In_img[y1,x2]
            f3 = In_img[y2,x1]
            f4 = In_img[y2,x2]

            interpolation = (
                f1 * (1 - x_weight) * (1 - y_weight) +
                f2 * x_weight * (1 - y_weight) +
                f3 * (1 - x_weight) * y_weight +
                f4 * x_weight * y_weight
            )

            Out_img[iY, iX] = interpolation

    Out_img = np.clip(Out_img, 0, 255).astype(np.uint8)

    return Out_img

# Shrinking methods
 
def MeanResizing(In_img,scale):
    Out_img=In_img
    if In_img.max() <= 1.0:
        In_img = In_img * 255.0

    height = In_img.shape[0]
    width = In_img.shape[1]

    new_width = int(np.round(width * scale))
    new_height = int(np.round(height * scale))

    if len(In_img.shape) == 3:
        Out_img = np.zeros((new_height, new_width, In_img.shape[2]), dtype=np.float32)
    else:
        Out_img = np.zeros((new_height, new_width), dtype=np.float32)

    Y_points = np.linspace(0, height - 1, new_height)
    X_points = np.linspace(0, width - 1, new_width)
            
    for iX, x_val in enumerate(X_points):
        for iY, y_val in enumerate(Y_points):
            
            ix = np.round(x_val) + np.arange(-3, 4)
            iy = np.round(y_val) + np.arange(-3, 4)

            ix = np.clip(ix, 0, width - 1).astype(int)
            iy = np.clip(iy,0,height - 1).astype(int)
            
            fragment = In_img[iy[0]:iy[-1]+1, ix[0]:ix[-1]+1]
            fragment_mean = np.mean(fragment, axis=(0,1))

            Out_img[iY, iX] = fragment_mean

      
    ####
    Out_img = np.clip(Out_img, 0, 255).astype(np.uint8)
    return Out_img



def WeightedMeanResizing(In_img, scale):
    height = In_img.shape[0]
    width = In_img.shape[1]
    if In_img.max() <= 1.0:
        In_img = In_img * 255.0

    new_width = int(np.round(width * scale))
    new_height = int(np.round(height * scale))

    if len(In_img.shape) == 3:
        Out_img = np.zeros((new_height, new_width, In_img.shape[2]), dtype=np.float32)
    else:
        Out_img = np.zeros((new_height, new_width), dtype=np.float32)

    Y_points = np.linspace(0, height - 1, new_height)
    X_points = np.linspace(0, width - 1, new_width)

    base_weights = np.array([
        [1, 2, 3, 4, 3, 2, 1],
        [2, 4, 6, 8, 6, 4, 2],
        [3, 6, 9, 12, 9, 6, 3],
        [4, 8, 12, 16, 12, 8, 4],
        [3, 6, 9, 12, 9, 6, 3],
        [2, 4, 6, 8, 6, 4, 2],
        [1, 2, 3, 4, 3, 2, 1]
    ], dtype=np.float32)
            
    for iX, x_val in enumerate(X_points):
        for iY, y_val in enumerate(Y_points):
            
            raw_ix = np.round(x_val) + np.arange(-3, 4)
            raw_iy = np.round(y_val) + np.arange(-3, 4)

            valid_x = (raw_ix >= 0) & (raw_ix < width)
            valid_y = (raw_iy >= 0) & (raw_iy < height)

            ix = raw_ix[valid_x].astype(int)
            iy = raw_iy[valid_y].astype(int)
            fragment = In_img[iy[0]:iy[-1]+1, ix[0]:ix[-1]+1]

            w = base_weights[valid_y, :][:, valid_x]

            if len(In_img.shape) == 3:
                w_3d = w[:, :, np.newaxis]

                fragment_mean = np.sum(fragment * w_3d, axis=(0, 1)) / np.sum(w)
            else:

                fragment_mean = np.sum(fragment * w) / np.sum(w)

            Out_img[iY, iX] = fragment_mean
            
    Out_img = np.clip(Out_img, 0, 255).astype(np.uint8)
    return Out_img

def MedianResizing(In_img,scale):
        Out_img=In_img
        if In_img.max() <= 1.0:
            In_img = In_img * 255.0

        height = In_img.shape[0]
        width = In_img.shape[1]

        new_width = int(np.round(width * scale))
        new_height = int(np.round(height * scale))

        if len(In_img.shape) == 3:
            Out_img = np.zeros((new_height, new_width, In_img.shape[2]), dtype=np.float32)
        else:
            Out_img = np.zeros((new_height, new_width), dtype=np.float32)

        Y_points = np.linspace(0, height - 1, new_height)
        X_points = np.linspace(0, width - 1, new_width)
                
        for iX, x_val in enumerate(X_points):
            for iY, y_val in enumerate(Y_points):
                
                ix = np.round(x_val) + np.arange(-3, 4)
                iy = np.round(y_val) + np.arange(-3, 4)

                ix = np.clip(ix, 0, width - 1).astype(int)
                iy = np.clip(iy,0,height - 1).astype(int)
                
                fragment = In_img[iy[0]:iy[-1]+1, ix[0]:ix[-1]+1]
                fragment_mean = np.median(fragment, axis=(0,1))

                Out_img[iY, iX] = fragment_mean
            
        Out_img = np.clip(Out_img, 0, 255).astype(np.uint8)
        return Out_img


def EdgeDetection(img):
    processed_img = img.copy()
    
    if processed_img.dtype != np.uint8:
        if processed_img.max() <= 1.0:
            processed_img = processed_img * 255
            
        processed_img = np.clip(processed_img, 0, 255).astype(np.uint8)

    edges = cv2.Canny(processed_img, 100, 200)
 

    return edges

##########################################
### Main Program  ########################
##########################################

def plot_resize(img, scale, nnscale, bscale, ed_img, ed_nnscale, ed_bscale, mr_img, wmr_img, mdr_img, ed_mr_img, ed_wmr_img,ed_mdr_img, oROI, filename,counter,figsize=(5,5)):
    ROI=(np.array(oROI)*scale).astype(int)
    f,axs=plt.subplots(4,3,num=counter,figsize=figsize) 
    f.suptitle(f"{filename} ROI: {ROI}")
    axs[0,0].imshow(img[oROI[1]:oROI[1]+oROI[3],oROI[0]:oROI[0]+oROI[2],:])
    axs[0,0].set_title("Original")
    axs[0,0].set_axis_off()

    axs[0,1].imshow(nnscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
    axs[0,1].set_title(f"NN scale {scale}")
    axs[0,1].set_axis_off()


    axs[0,2].imshow(bscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
    axs[0,2].set_title(f"Blinear scale {scale}")
    axs[0,2].set_axis_off()

    if len(ed_img.shape)==3:
        axs[1,0].imshow(ed_img[oROI[1]:oROI[1]+oROI[3],oROI[0]:oROI[0]+oROI[2],:])
        axs[1,0].set_title("Edges Original")
        axs[1,0].set_axis_off()
    else:
        axs[1,0].imshow(ed_img[oROI[1]:oROI[1]+oROI[3],oROI[0]:oROI[0]+oROI[2]])
        axs[1,0].set_title("Edges Original")
        axs[1,0].set_axis_off()

    if len(ed_nnscale.shape)==3:
        axs[1,1].imshow(ed_nnscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
        axs[1,1].set_title("Edges NN")
        axs[1,1].set_axis_off()
    else:
        axs[1,1].imshow(ed_nnscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2]])
        axs[1,1].set_title("Edges NN")
        axs[1,1].set_axis_off()
        
    if len(ed_bscale.shape)==3:
        axs[1,2].imshow(ed_bscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
        axs[1,2].set_title("Edges Bilinear")
        axs[1,2].set_axis_off()
    else:
        axs[1,2].imshow(ed_bscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2]])
        axs[1,2].set_title("Edges Bilinear")
        axs[1,2].set_axis_off()

    axs[2,0].imshow(mr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
    axs[2,0].set_title(f"Mean Resizing scale {scale}")
    axs[2,0].set_axis_off()

    axs[2,1].imshow(wmr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
    axs[2,1].set_title(f"Weighted Mean Resizing scale {scale}")
    axs[2,1].set_axis_off()

    axs[2,2].imshow(mdr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
    axs[2,2].set_title(f"Median Resizing scale {scale}")
    axs[2,2].set_axis_off()
        
    if len(ed_mr_img.shape)==3:
        axs[3,0].imshow(ed_mr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
        axs[3,0].set_title("Edges Mean Resizing")
        axs[3,0].set_axis_off()
    else:
        axs[3,0].imshow(ed_mr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2]])
        axs[3,0].set_title("Edges Mean Resizing")
        axs[3,0].set_axis_off()
        
    if len(ed_wmr_img.shape)==3:
        axs[3,1].imshow(ed_wmr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
        axs[3,1].set_title("Edges Weighted Mean Resizing")
        axs[3,1].set_axis_off()
    else:
        # axs[3,1].imshow(ed_nnscale[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2]])
        axs[3,1].imshow(ed_wmr_img[ROI[1]:ROI[1]+ROI[3], ROI[0]:ROI[0]+ROI[2]])

        axs[3,1].set_title("Edges Weighted Mean Resizing")
        axs[3,1].set_axis_off()
        
    if len(ed_mdr_img.shape)==3:
        axs[3,2].imshow(ed_mdr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2],:])
        axs[3,2].set_title("Edges Median Resizing")
        axs[3,2].set_axis_off()
    else:
        axs[3,2].imshow(ed_mdr_img[ROI[1]:ROI[1]+ROI[3],ROI[0]:ROI[0]+ROI[2]])
        axs[3,2].set_title("Edges Median Resizing")
        axs[3,2].set_axis_off()
        
    return f

def plot_scaling(img, scale, nnscale, bscale, counter, ed_img, ed_nnscale, ed_bscale, file,figsize=(5,5)):
    f,axs=plt.subplots(2,3,num=counter,figsize=figsize) 
    f.suptitle(f"{file}")

    axs[0,0].imshow(img)
    axs[0,0].set_title("Original")
    axs[0,0].set_axis_off()

    axs[0,1].imshow(nnscale)
    axs[0,1].set_title(f"NN scale {scale}")
    axs[0,1].set_axis_off()

    axs[0,2].imshow(bscale)
    axs[0,2].set_title(f"Blinear scale {scale}")
    axs[0,2].set_axis_off()

    axs[1,0].imshow(ed_img)
    axs[1,0].set_title("Edges Original")
    axs[1,0].set_axis_off()

    axs[1,1].imshow(ed_nnscale)
    axs[1,1].set_title("Edges NN")
    axs[1,1].set_axis_off()

    axs[1,2].imshow(ed_bscale)
    axs[1,2].set_title("Edges Bilinear")
    axs[1,2].set_axis_off()
    return f

if Test:
    # test case
    if Scaling_test:
        img= np.zeros((3,3,3),dtype=np.float32)
        img[1,1,:]=1.0
        for scale in ScalesUp:
            f,axs=plt.subplots(1,2)
            nnscale=NearestNeigbourScaling(img,scale)
            axs[0].imshow(nnscale)
            bscale=BilinearScaling(img,scale)
            axs[1].imshow(bscale)

    else:
        counter=1
        for scale in ScalesUp:
            img=plt.imread(os.path.join(ImgDir,SmallImages[0]))
            nnscale=NearestNeigbourScaling(img,scale)
            bscale=BilinearScaling(img,scale)
            ed_img=EdgeDetection(img)
            ed_nnscale=EdgeDetection(nnscale)
            ed_bscale=EdgeDetection(bscale)

            f = plot_scaling(img, scale, nnscale, bscale, counter, ed_img, ed_nnscale, ed_bscale, SmallImages[0])

            counter+=1

        for scale in ScalesDown:    
            img=plt.imread(os.path.join(ImgDir,BigImages[0]["Filename"]))

            nnscale=NearestNeigbourScaling(img,scale)
            bscale=BilinearScaling(img,scale)

            ed_img=EdgeDetection(img)
            ed_nnscale=EdgeDetection(nnscale)
            ed_bscale=EdgeDetection(bscale)

            mr_img=MeanResizing(img,scale)
            wmr_img=WeightedMeanResizing(img,scale)
            mdr_img=MedianResizing(img,scale)

            ed_mr_img=EdgeDetection(mr_img)
            ed_wmr_img=EdgeDetection(wmr_img)
            ed_mdr_img=EdgeDetection(mdr_img)
  
            f= plot_resize(img, scale, nnscale, bscale, ed_img, ed_nnscale, ed_bscale, mr_img, wmr_img, mdr_img, ed_mr_img, ed_wmr_img,ed_mdr_img, BigImages[0]["ROIs"][0], BigImages[0]["Filename"],counter=counter)
            counter+=1
        
    plt.savefig('test.png')

else: 
    # generate raport
    document = Document()
    document.add_heading('Report',0) # tworzenie nagłówków druga wartość to poziom nagłówka 
    document.add_paragraph("Autor: ")
    document.add_paragraph("Proszę wstawić mi 2 jeżeli tego nie wyedytuję")
    document.add_section()
    document.add_heading("Test algorytmów powiększania",1)
    counter = 1 
    for file in SmallImages:
        img=plt.imread(os.path.join(ImgDir,file))
        for scale in ScalesUp:
            nnscale=NearestNeigbourScaling(img,scale)
            bscale=BilinearScaling(img,scale)
            ed_img=EdgeDetection(img)
            ed_nnscale=EdgeDetection(nnscale)
            ed_bscale=EdgeDetection(bscale)
            
            f = plot_scaling(img, scale, nnscale, bscale, counter, ed_img, ed_nnscale, ed_bscale, file) # set figszie

            memfile = BytesIO() 
            f.savefig(memfile)
            document.add_picture(memfile, width=Inches(6)) # set document size
            memfile.close()
            plt.close(f)
            counter += 1
    document.add_section()
    document.add_heading("Test algorytmów pomniejszania",1)
    for file_dict in BigImages:
        filename=file_dict['Filename']
        img=plt.imread(os.path.join(ImgDir,filename))
        for scale in ScalesDown:
            nnscale=NearestNeigbourScaling(img,scale)
            bscale=BilinearScaling(img,scale)

            ed_img=EdgeDetection(img)
            ed_nnscale=EdgeDetection(nnscale)
            ed_bscale=EdgeDetection(bscale)

            mr_img=MeanResizing(img,scale)
            wmr_img=WeightedMeanResizing(img,scale)
            mdr_img=MedianResizing(img,scale)

            ed_mr_img=EdgeDetection(mr_img)
            ed_wmr_img=EdgeDetection(wmr_img)
            ed_mdr_img=EdgeDetection(mdr_img)

            for ROI in file_dict['ROIs']:

                f = plot_resize(img, scale, nnscale, bscale, ed_img, ed_nnscale, ed_bscale, mr_img, wmr_img, mdr_img, ed_mr_img, ed_wmr_img,ed_mdr_img, ROI, filename,counter=counter) # set figszie

                memfile = BytesIO() 
                f.savefig(memfile)
                document.add_picture(memfile, width=Inches(6)) 
                memfile.close()
                plt.close(f)
                counter += 1
    document.add_section()
    document.add_heading("Podsumowanie i wnioski",1)
    document.add_paragraph("Tu proszę zebrać wszystkie obserwacje na podstawie powyższych wykresów i napisać wnioski.")
    document.add_paragraph("Powiększanie: W przypadku powiększania metoda najbliższych sąsiadów działa lepiej od metody bilinearnej. Przy powiększeniu w skali 2 obydwie metody działają podobnie i, patrząc na detekcję krawędzi, nawet w przypadku zdjęcia „IMG_SMALL/SMALL_0001.tif” metoda bilinearna wykazała bardziej szczegółowy obraz, lecz w innych skalach powiększania (3, 5) metoda najbliższego sąsiada daje dużo bardziej widoczne i ostre krawędzie. Jednak jakość obrazków wciąż pozostaje na podobnym poziomie, porównywalnym do oryginału, ponieważ na wykresach różnice dokładnie widać dopiero po analizie krawędzi.")
    document.add_paragraph("Pomniejszanie: W przypadku pomniejszania wszystkie metody prezentują podobną skuteczność, to znaczy detekcja krawędzi we wszystkich metodach daje podobne rezultaty, lecz w przypadku pomniejszania jakość zdjęć znacząco spada — zdjęcia stają się dużo mniej wyraźne, a piksele stają się bardziej ujednolicone i mniej szczegółowe.")

    
    document.save(OutputRaportFile) 
