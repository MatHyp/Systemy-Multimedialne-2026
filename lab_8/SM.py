import numpy as np
import matplotlib.pyplot as plt
import cv2
import os
from docx import Document
from docx.shared import Inches
from io import BytesIO
import scipy.fftpack

##########################################
### Settings #############################
##########################################

Test = False

OutputRaportFile = "yes.docx" 

Chroma_options = ["4:4:4", "4:2:2"]
Quant_options = [True, False]

QY = np.array([
        [16, 11, 10, 16, 24,  40,  51,  61],
        [12, 12, 14, 19, 26,  58,  60,  55],
        [14, 13, 16, 24, 40,  57,  69,  56],
        [14, 17, 22, 29, 51,  87,  80,  62],
        [18, 22, 37, 56, 68,  109, 103, 77],
        [24, 36, 55, 64, 81,  104, 113, 92],
        [49, 64, 78, 87, 103, 121, 120, 101],
        [72, 92, 95, 98, 112, 100, 103, 99],
        ])

QC = np.array([
        [17, 18, 24, 47, 99, 99, 99, 99],
        [18, 21, 26, 66, 99, 99, 99, 99],
        [24, 26, 56, 99, 99, 99, 99, 99],
        [47, 66, 99, 99, 99, 99, 99, 99],
        [99, 99, 99, 99, 99, 99, 99, 99],
        [99, 99, 99, 99, 99, 99, 99, 99],
        [99, 99, 99, 99, 99, 99, 99, 99],
        [99, 99, 99, 99, 99, 99, 99, 99],
        ])

QN = np.ones((8,8))

##########################################
### Data Set #############################
##########################################

ImgDir = r'.' # Address of folder with files (do nor delete `r``)

Images = [ #list of dictionaries
    {
        "Filename":"img.jpg", # File name (podmień na własny)
        "ROIs":[[0, 0, 128, 128]] # list of Region of interests
    }
]


##########################################
### Classes & Helpers ####################
##########################################

class ver2:
    def __init__(self, Y, Cb, Cr, OGShape, Ratio="4:4:4", QY=np.ones((8,8)), QC=np.ones((8,8))):
        self.shape = OGShape
        self.Y = Y
        self.Cb = Cb
        self.Cr = Cr
        self.ChromaRatio = Ratio
        self.QY = QY
        self.QC = QC

JPEG_class = ver2

def pad_image(img, block_size=16):
    h, w = img.shape[:2]
    pad_h = (block_size - (h % block_size)) % block_size
    pad_w = (block_size - (w % block_size)) % block_size
    
    if pad_h > 0 or pad_w > 0:
        img = np.pad(img, ((0, pad_h), (0, pad_w), (0, 0)), mode='edge')
    return img, (h, w)

def restore_dynamic_range(layer, orig_min, orig_max):
    curr_min, curr_max = layer.min(), layer.max()
    if curr_max - curr_min == 0:
        return layer
    normalized = (layer - curr_min) / (curr_max - curr_min)
    return normalized * (orig_max - orig_min) + orig_min


# byte run 

def byterun_count_repeats(data, start_idx):
    count = 1
    n = len(data)
    for i in range(start_idx, n - 1):
        if data[i] == data[i+1]:
            count += 1
        else:
            break
    return count

def byterun_count_differences(data, start_idx):
    count = 1
    n = len(data)
    for i in range(start_idx + 1, n):
        if i < n - 1 and data[i] == data[i+1]:
            break
        count += 1
    return count

def byterun_encoder(image):
    original_shape = np.array(image).shape
    shape_info = np.array([len(original_shape)] + list(original_shape), dtype=int)
    data = np.array(image).flatten().astype(int)

    if len(data) == 0:
        return shape_info

    result = np.zeros(data.shape[0] * 2, dtype=int)
    write_index = 0
    i = 0
    n = len(data)
    
    while i < n:
        if i < n - 1 and data[i] == data[i+1]:
            total_count = byterun_count_repeats(data, i)
            counter = total_count
            while counter > 128:
                result[write_index] = -127        
                result[write_index + 1] = data[i] 
                write_index += 2
                counter -= 128            
            if counter > 0:
                result[write_index] = -(counter - 1)
                result[write_index + 1] = data[i]
                write_index += 2
            i += total_count
        else:
            total_count = byterun_count_differences(data, i)
            counter = total_count
            processed = 0
            while counter > 128:
                result[write_index] = 127
                write_index += 1
                result[write_index : write_index + 128] = data[i + processed : i + processed + 128]
                write_index += 128
                processed += 128
                counter -= 128
            if counter > 0:
                result[write_index] = counter - 1
                write_index += 1
                result[write_index : write_index + counter] = data[i + processed : i + processed + counter]
                write_index += counter
            i += total_count

    compressed_data = result[:write_index]
    return np.concatenate((shape_info, compressed_data))

def byterun_decoder(encoded_data):
    dims = encoded_data[0]
    shape = tuple(encoded_data[1 : 1 + dims])
    data = encoded_data[1 + dims :]
    
    total_elements = np.prod(shape)
    result = np.zeros(total_elements, dtype=int)
    
    i = 0
    write_index = 0
    n = len(data)
    
    while i < n:
        flag = data[i]
        if flag < 0:
            count = -flag + 1
            sym = data[i+1]
            result[write_index : write_index + count] = sym
            write_index += count
            i += 2
        else: 
            count = flag + 1
            result[write_index : write_index + count] = data[i+1 : i+1+count]
            write_index += count
            i += 1 + count
            
    return result.reshape(shape)

##########################################
### JPEG Core Functions ##################
##########################################

def dct2(a):
    return scipy.fftpack.dct(scipy.fftpack.dct(a.astype(float), axis=0, norm='ortho'), axis=1, norm='ortho')

def idct2(a):
    return scipy.fftpack.idct(scipy.fftpack.idct(a.astype(float), axis=0, norm='ortho'), axis=1, norm='ortho')

def zigzag(A):
    template = np.array([
        [ 0,  1,  5,  6, 14, 15, 27, 28],
        [ 2,  4,  7, 13, 16, 26, 29, 42],
        [ 3,  8, 12, 17, 25, 30, 41, 43],
        [ 9, 11, 18, 24, 31, 40, 44, 53],
        [10, 19, 23, 32, 39, 45, 52, 54],
        [20, 22, 33, 38, 46, 51, 55, 60],
        [21, 34, 37, 47, 50, 56, 59, 61],
        [35, 36, 48, 49, 57, 58, 62, 63]
    ])
    
    if len(A.shape) == 1:
        B = np.zeros((8, 8))
        for r in range(0, 8):
            for c in range(0, 8):
                B[r, c] = A[template[r, c]]
    else:
        B = np.zeros((64,))
        for r in range(0, 8):
            for c in range(0, 8):
                B[template[r, c]] = A[r, c]
    return B

def CompressBlock(block, Q):
    centered = block.astype(float) - 128
    dct = dct2(centered)
    quantized = np.round(dct / Q).astype(int)
    return zigzag(quantized)

def DecompressBlock(vector, Q):
    de_zigzaged = zigzag(vector)
    de_quantized = de_zigzaged * Q
    idct = idct2(de_quantized)
    return idct + 128

def CompressLayer(L, Q):
    S = np.array([])
    for w in range(0, L.shape[0], 8):
        for k in range(0, L.shape[1], 8):
            block = L[w:(w+8), k:(k+8)]
            S = np.append(S, CompressBlock(block, Q))
    return S

def DecompressLayer(S, Q, shape):
    L = np.zeros(shape)
    m = shape[1] / 8
    for idx, i in enumerate(range(0, S.shape[0], 64)):
        vector = S[i:(i+64)]
        k = int((idx % m) * 8)
        w = int((idx // m) * 8)
        L[w:(w+8), k:(k+8)] = DecompressBlock(vector, Q)
    return L

def CompressJPEG(RGB, Ratio="4:4:4", QY=np.ones((8,8)), QC=np.ones((8,8))):
    padded_RGB, original_dims = pad_image(RGB, block_size=16)
    YCrCb = cv2.cvtColor(padded_RGB.astype(np.uint8), cv2.COLOR_RGB2YCrCb).astype(int)
    
    Y = YCrCb[:, :, 0]
    Cr = YCrCb[:, :, 1]
    Cb = YCrCb[:, :, 2]
    
    min_max = {
        'Y': (Y.min(), Y.max()),
        'Cr': (Cr.min(), Cr.max()),
        'Cb': (Cb.min(), Cb.max())
    }
    
    if Ratio == "4:2:2":
        Cr = Cr[:, ::2]
        Cb = Cb[:, ::2]
        
    Y_comp = CompressLayer(Y, QY)
    Cb_comp = CompressLayer(Cb, QC)
    Cr_comp = CompressLayer(Cr, QC)
    
    # Kompresja ByteRun
    Y_encoded = byterun_encoder(Y_comp)
    Cb_encoded = byterun_encoder(Cb_comp)
    Cr_encoded = byterun_encoder(Cr_comp)
    
    stats = {
        'Y_orig': len(Y_comp), 'Y_comp': len(Y_encoded),
        'Cb_orig': len(Cb_comp), 'Cb_comp': len(Cb_encoded),
        'Cr_orig': len(Cr_comp), 'Cr_comp': len(Cr_encoded)
    }
        
    JPEG = JPEG_class(
        Y = Y_encoded,
        Cb = Cb_encoded,
        Cr = Cr_encoded,
        OGShape = padded_RGB.shape,
        Ratio = Ratio,
        QY = QY,
        QC = QC
    )
    
    JPEG.original_dims = original_dims
    JPEG.min_max = min_max
    JPEG.stats = stats
    
    return JPEG

def DecompressJPEG(JPEG):
    # Dekompresja ByteRun
    Y_decoded = byterun_decoder(JPEG.Y)
    Cb_decoded = byterun_decoder(JPEG.Cb)
    Cr_decoded = byterun_decoder(JPEG.Cr)

    shape_Y = (JPEG.shape[0], JPEG.shape[1])
    Y = DecompressLayer(Y_decoded, JPEG.QY, shape_Y)
    
    if JPEG.ChromaRatio == "4:2:2":
        shape_C = (JPEG.shape[0], JPEG.shape[1] // 2)
    else:
        shape_C = shape_Y
        
    Cr = DecompressLayer(Cr_decoded, JPEG.QC, shape_C)
    Cb = DecompressLayer(Cb_decoded, JPEG.QC, shape_C)
    
    if JPEG.ChromaRatio == "4:2:2":
        Cr = np.repeat(Cr, 2, axis=1)
        Cb = np.repeat(Cb, 2, axis=1)
        
    if hasattr(JPEG, 'min_max'):
        Y = restore_dynamic_range(Y, JPEG.min_max['Y'][0], JPEG.min_max['Y'][1])
        Cr = restore_dynamic_range(Cr, JPEG.min_max['Cr'][0], JPEG.min_max['Cr'][1])
        Cb = restore_dynamic_range(Cb, JPEG.min_max['Cb'][0], JPEG.min_max['Cb'][1])
        
    YCrCb = np.dstack([Y, Cr, Cb]).clip(0, 255).astype(np.uint8)
    RGB = cv2.cvtColor(YCrCb, cv2.COLOR_YCrCb2RGB)
    
    if hasattr(JPEG, 'original_dims'):
        orig_h, orig_w = JPEG.original_dims
        RGB = RGB[:orig_h, :orig_w]
        
    return RGB

##########################################
### Main Program #########################
##########################################

def plot_comparisone(counter, OG, Decomp, figsize=(5,8)):
    fig, axs = plt.subplots(4, 2, num=counter, sharex=True, sharey=True, figsize=figsize)
    
    axs[0,0].imshow(OG) 
    PRZED_YCrCb = cv2.cvtColor(OG, cv2.COLOR_RGB2YCrCb)
    axs[1,0].imshow(PRZED_YCrCb[:,:,0], cmap='gray') 
    axs[2,0].imshow(PRZED_YCrCb[:,:,1], cmap='gray')
    axs[3,0].imshow(PRZED_YCrCb[:,:,2], cmap='gray')
    
    axs[0,0].set_title("Oryginał")
    axs[1,0].set_title("Y")
    axs[2,0].set_title("Cr")
    axs[3,0].set_title("Cb")

    axs[0,1].imshow(Decomp) 
    PO_YCrCb = cv2.cvtColor(Decomp, cv2.COLOR_RGB2YCrCb)
    axs[1,1].imshow(PO_YCrCb[:,:,0], cmap='gray')
    axs[2,1].imshow(PO_YCrCb[:,:,1], cmap='gray')
    axs[3,1].imshow(PO_YCrCb[:,:,2], cmap='gray')
    
    axs[0,1].set_title("Po kompresji")
    axs[1,1].set_title("Y")
    axs[2,1].set_title("Cr")
    axs[3,1].set_title("Cb")
    
    for ax in axs.flatten():
        ax.set_axis_off() 
    
    return fig

if Test:
    img = plt.imread(os.path.join(ImgDir, Images[0]["Filename"]))
    ROI = Images[0]["ROIs"][0]
    fragment = img[ROI[1]:ROI[1]+ROI[3], ROI[0]:ROI[0]+ROI[2]]
    Counter = 1
    
    for Chroma in Chroma_options:
        for Quant in Quant_options:
            if Quant:
                tJPEG = CompressJPEG(fragment, Ratio=Chroma, QY=QY, QC=QC)
            else:
                tJPEG = CompressJPEG(fragment, Ratio=Chroma, QY=QN, QC=QN)
                
            New_Fragment = DecompressJPEG(tJPEG)
                
            f = plot_comparisone(Counter, fragment, New_Fragment)
            f.suptitle(f"Plik: {Images[0]['Filename']} Chroma: {Chroma} Kwantyzacja: {Quant}")
            
            # Wypisywanie statystyk w konsoli
            print(f"\n--- Statystyki ByteRun (Chroma: {Chroma}, Kwantyzacja: {Quant}) ---")
            pY = (tJPEG.stats['Y_comp'] / tJPEG.stats['Y_orig']) * 100
            pCb = (tJPEG.stats['Cb_comp'] / tJPEG.stats['Cb_orig']) * 100
            pCr = (tJPEG.stats['Cr_comp'] / tJPEG.stats['Cr_orig']) * 100
            print(f"Y: {pY:.2f}%, Cb: {pCb:.2f}%, Cr: {pCr:.2f}% (oryginał vs skompresowany)")

            Counter += 1
    plt.show()
    
else:
    document = Document()
    document.add_heading('Raport z działania kompresji JPEG', 0) 
    document.add_paragraph("Autor: Mateusz Hypś")
    document.add_section()
    document.add_heading("Fragmenty wygenerowane na podstawie działania funkcji", 1)
    
    Counter = 1
    for file_dict in Images:
        filename = file_dict['Filename']
        img_path = os.path.join(ImgDir, filename)
        
        if not os.path.exists(img_path):
            print(f"Brak pliku: {img_path}. Pomijanie...")
            continue
            
        img = plt.imread(img_path)
        
        for ROI in file_dict['ROIs']:
            fragment = img[ROI[1]:ROI[1]+ROI[3], ROI[0]:ROI[0]+ROI[2]]
            
            for Chroma in Chroma_options:
                for Quant in Quant_options:
                    if Quant:
                        tJPEG = CompressJPEG(fragment, Ratio=Chroma, QY=QY, QC=QC)
                    else:
                        tJPEG = CompressJPEG(fragment, Ratio=Chroma, QY=QN, QC=QN)
                        
                    New_Fragment = DecompressJPEG(tJPEG)
                        
                    f = plot_comparisone(Counter, fragment, New_Fragment)
                    f.suptitle(f"Plik: {filename} Chroma: {Chroma} Kwantyzacja: {Quant}")
                    
                    memfile = BytesIO() 
                    f.savefig(memfile)
                    document.add_picture(memfile, width=Inches(6))
                    memfile.close()
                    f.clf()
                    
                    # Dodawanie statystyk ByteRun do dokumentu DOCX
                    pY = (tJPEG.stats['Y_comp'] / tJPEG.stats['Y_orig']) * 100
                    pCb = (tJPEG.stats['Cb_comp'] / tJPEG.stats['Cb_orig']) * 100
                    pCr = (tJPEG.stats['Cr_comp'] / tJPEG.stats['Cr_orig']) * 100
                    
                    document.add_paragraph(f"Wyniki kompresji bezstratnej ByteRun:")
                    document.add_paragraph(f"- Warstwa Y wektor skrócony do: {pY:.2f}% oryginalnej długości")
                    document.add_paragraph(f"- Warstwa Cb wektor skrócony do: {pCb:.2f}% oryginalnej długości")
                    document.add_paragraph(f"- Warstwa Cr wektor skrócony do: {pCr:.2f}% oryginalnej długości")
                    document.add_paragraph("") # Pusty odstęp pomiędzy obrazkami
                    
    document.add_section()
    document.add_heading("Podsumowanie i wnioski", 1)
    document.add_paragraph("Tu proszę zebrać wszystkie obserwacje na podstawie powyższych wykresów i napisać wnioski.")
    document.save(OutputRaportFile)
    print(f"Zapisano pomyślnie do pliku {OutputRaportFile}")