import os
from docx import Document
from docx.shared import Inches

# Inicjalizacja dokumentu
document = Document()

# --- TYTUŁ I AUTOR ---
document.add_heading('Sprawozdanie: Kompresja Wideo', 0)
document.add_paragraph("Autor: Mateusz Hypś")
document.add_paragraph("Plik testowy: clip_1.mp4")

# ==============================================================================
# CZĘŚĆ 1: Badanie jakości
# ==============================================================================
document.add_heading("Część 1: Badanie jakości dla różnych parametrów (bez RLE)", 1)
document.add_paragraph(
    "Celem tej części było zbadanie wpływu redukcji chrominancji (Chroma Subsampling) oraz "
    "dzielnika przy kodowaniu różnicowym na jakość obrazu. Poniżej przedstawiono obszerną analizę "
    "wizualną na rozbitych warstwach Luminancji (Y) i Chrominancji (Cb, Cr) dla różnych klatek wideo, "
    "aby zaobserwować wpływ kompresji w różnych fazach ruchu obiektu."
)

# Zestawienie aż 8 zdjęć z klatek 11, 15 i 19
zdjecia_czesc1 = [
    # KLATKA 11
    {
        "plik": "wyniki_jakosc_2/clip_1_roi0_klatka11_sub_4-2-2_dzielnik_4_RLE_OFF.png",
        "tytul": "Klatka 11 - Lekka kompresja (Subsampling 4:2:2, Dzielnik 4)",
        "opis": "Początkowa faza ruchu. Subtelna redukcja rozdzielczości kolorów w poziomie. Różnice na warstwach Cb i Cr są minimalne."
    },
    {
        "plik": "wyniki_jakosc_2/clip_1_roi0_klatka11_sub_4-1-0_dzielnik_4_RLE_OFF.png",
        "tytul": "Klatka 11 - Bardzo mocna kompresja (Subsampling 4:1:0, Dzielnik 4)",
        "opis": "Dla porównania skrajny przypadek. Widać wyraźne rozmycie na warstwach chrominancji już w początkowej fazie ruchu."
    },
    
    # KLATKA 15
    {
        "plik": "wyniki_jakosc_2/clip_1_roi0_klatka15_sub_4-2-0_dzielnik_4_RLE_OFF.png",
        "tytul": "Klatka 15 - Umiarkowana kompresja (Subsampling 4:2:0, Dzielnik 4)",
        "opis": "Środkowa faza ruchu. Próbkowanie 4:2:0 radzi sobie dobrze, pomimo widocznego przemieszczenia obiektu."
    },
    {
        "plik": "wyniki_jakosc_2/clip_1_roi0_klatka15_sub_4-1-1_dzielnik_4_RLE_OFF.png",
        "tytul": "Klatka 15 - Mocna kompresja (Subsampling 4:1:1, Dzielnik 4)",
        "opis": "Agresywne cięcie informacji w poziomie zaczyna generować zauważalne błędy na krawędziach poruszającego się obiektu."
    },

    # KLATKA 19 (Pełne zestawienie, podsumowanie)
    {
        "plik": "wyniki_jakosc_2/clip_1_roi0_klatka19_sub_4-2-2_dzielnik_4_RLE_OFF.png",
        "tytul": "Klatka 19 - Subsampling 4:2:2, Dzielnik 4",
        "opis": "Końcowa badana faza ruchu. Kompresja 4:2:2 wciąż zachowuje bardzo dobrą jakość szczegółów barwnych."
    },
    {
        "plik": "wyniki_jakosc_2/clip_1_roi0_klatka19_sub_4-2-0_dzielnik_4_RLE_OFF.png",
        "tytul": "Klatka 19 - Złoty środek (Subsampling 4:2:0, Dzielnik 4)",
        "opis": "Kolor próbkowany rzadziej w obu kierunkach. Optymalny balans między utratą jakości a oszczędnością miejsca."
    },
    {
        "plik": "wyniki_jakosc_2/clip_1_roi0_klatka19_sub_4-1-1_dzielnik_4_RLE_OFF.png",
        "tytul": "Klatka 19 - Subsampling 4:1:1, Dzielnik 4",
        "opis": "Przy dynamicznym ruchu w klatce 19 warstwy Cb i Cr wykazują już widoczną, blokową strukturę błędów."
    },
    {
        "plik": "wyniki_jakosc_2/clip_1_roi0_klatka19_sub_4-1-0_dzielnik_4_RLE_OFF.png",
        "tytul": "Klatka 19 - Subsampling 4:1:0, Dzielnik 4",
        "opis": "Najwyższy stopień kompresji barwy. Obecne są rażące artefakty kompresyjne i wyraźne rozmycie krawędzi."
    }
]

# Wklejanie zdjęć do Worda
for zdjecie in zdjecia_czesc1:
    document.add_heading(zdjecie['tytul'], 2)
    document.add_paragraph(zdjecie['opis'])
    if os.path.exists(zdjecie['plik']):
        document.add_picture(zdjecie['plik'], width=Inches(6.0))
    else:
        document.add_paragraph(f"[BŁĄD: Nie znaleziono pliku {zdjecie['plik']}]")

# Wniosek do części 1
document.add_paragraph(
    "Wniosek końcowy: Przeprowadzona szeroka analiza na różnych klatkach dowodzi, że ustawienia Subsampling 4:2:0 "
    "oraz Dzielnik 4 stanowią najbardziej optymalny wybór. Gwarantują one wysoką redukcję ilości danych, zachowując "
    "obraz wolny od mocno rażących artefaktów, nawet podczas dynamicznego ruchu."
)

# ==============================================================================
# CZĘŚĆ 2: Wykresy pamięci
# ==============================================================================
document.add_page_break()
document.add_heading("Część 2: Badanie skuteczności kompresji z użyciem RLE", 1)
document.add_paragraph(
    "W oparciu o wybrane w Części 1 parametry (4:2:0, dzielnik 4), zastosowano bezstratną "
    "kompresję strumieniową RLE. Poniżej zbadano, jak interwał klatek kluczowych wpływa na "
    "zysk z kompresji warstw."
)

document.add_heading("Wykres referencyjny (BEZ użycia RLE)", 2)
document.add_paragraph(
    "Referencyjny przypadek udowadniający, że bez algorytmu pakującego (RLE), samo "
    "kodowanie różnicowe nie przynosi zysków pojemnościowych (zysk bliski 0%)."
)

plik_referencyjny = "wyniki_pamiec_2/wykres_clip_1_klatki_8_sub_4-2-0_dz_4_bez_RLE.png"
if os.path.exists(plik_referencyjny):
    document.add_picture(plik_referencyjny, width=Inches(6.0))
else:
    document.add_paragraph(f"[BŁĄD: Nie znaleziono pliku {plik_referencyjny}]")

document.add_heading("Wykresy z użyciem algorytmu RLE", 2)
document.add_paragraph("Zestawienie wyników przy włączonym algorytmie RLE dla rosnących odległości klatek kluczowych:")

wykresy_rle = [
    "wyniki_pamiec_2/wykres_clip_1_klatki_2_sub_4-2-0_dz_4_z_RLE.png",
    "wyniki_pamiec_2/wykres_clip_1_klatki_5_sub_4-2-0_dz_4_z_RLE.png",
    "wyniki_pamiec_2/wykres_clip_1_klatki_8_sub_4-2-0_dz_4_z_RLE.png",
    "wyniki_pamiec_2/wykres_clip_1_klatki_12_sub_4-2-0_dz_4_z_RLE.png",
    "wyniki_pamiec_2/wykres_clip_1_klatki_16_sub_4-2-0_dz_4_z_RLE.png"
]

for wykres in wykresy_rle:
    if os.path.exists(wykres):
        document.add_picture(wykres, width=Inches(6.0))
    else:
        document.add_paragraph(f"[BŁĄD: Nie znaleziono pliku {wykres}]")

# Zapis dokumentu
document.save('Sprawozdanie_Gotowe_Mateusz_v3.docx')
print("Pomyślnie wygenerowano rozbudowane sprawozdanie: Sprawozdanie_Gotowe_Mateusz_v3.docx")